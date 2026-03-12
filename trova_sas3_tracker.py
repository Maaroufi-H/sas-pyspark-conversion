"""
trova_sas3_tracker.py
=====================
Parser SAS con tracciatura gerarchica dei blocchi, macro-categorizzazione,
e classificazione automatica convertibile SI/NO con affidabilita'.

Genera: DataFrame pandas, Excel (.xlsx), JSON ad albero.
"""
from __future__ import annotations
from pathlib import Path
from typing import Dict, List, Optional, Tuple
from collections import deque
from datetime import datetime
import re, hashlib, json, argparse

import pandas as pd

# ═══════════════════════════════════════════════════════════════════
# IMPORT dal parser esistente (trova_sas2_modifica_run_macro.py)
# ═══════════════════════════════════════════════════════════════════
from trova_sas2_modifica_run_macro import (
    RE_COMMENT_LINE,
    RE_IFDO_START,
    RE_ELSEDO_START,
    RE_MACRO_ENDTOK,
    RE_MACRO_START,
    RE_MEND_TOKEN,
    BLOCKS,
    _key_data,
    _key_proc_sort,
    _key_proc_sql,
    _key_proc_append,
    _key_proc_dataset,
    _key_include,
    _key_macro,
    _key_if,
    _key_else,
    _key_runmacro,
    preprocess_lines_keep_comments,
    _stable_uid,
)

# ═══════════════════════════════════════════════════════════════════
# MACRO-CATEGORIZZAZIONE
# ═══════════════════════════════════════════════════════════════════
MACRO_CATEGORIE: Dict[str, str] = {
    "data":          "DATA_STEP",
    "proc_sql":      "PROC_SQL",
    "proc_sort":     "PROC_SORT",
    "proc_append":   "PROC_APPEND",
    "proc_dataset":  "PROC_DATASETS",
    "macro":         "MACRO_DEF",
    "ifdo":          "IF_BLOCK",
    "elsedo":        "ELSE_BLOCK",
    "include":       "INCLUDE",
    "run_macro":     "MACRO_CALL",
}

# ═══════════════════════════════════════════════════════════════════
# COLORI ANSI PER PARENT (console + file .txt)
# ═══════════════════════════════════════════════════════════════════
_ANSI_PALETTE: List[str] = [
    "\033[91m",  # rosso brillante
    "\033[92m",  # verde brillante
    "\033[93m",  # giallo brillante
    "\033[94m",  # blu brillante
    "\033[95m",  # magenta brillante
    "\033[96m",  # ciano brillante
    "\033[33m",  # giallo scuro
    "\033[35m",  # magenta scuro
    "\033[36m",  # ciano scuro
    "\033[32m",  # verde scuro
]
_ANSI_BOLD  = "\033[1m"
_ANSI_RESET = "\033[0m"
_ANSI_ROOT  = "\033[97m"   # bianco brillante per blocchi ROOT (senza parent)


# ═══════════════════════════════════════════════════════════════════
# CLASSIFICATORE AUTOMATICO — MOTORE DINAMICO da file JSON esterno
# ═══════════════════════════════════════════════════════════════════
# Percorso default del file regole (stesso directory di questo script)
_DEFAULT_RULES_PATH = Path(__file__).parent / "regole_classificazione.json"

# Cache globale: caricata una sola volta
_REGOLE_CACHE: Dict[str, list] | None = None


def carica_regole(rules_path: Path | None = None) -> Dict[str, list]:
    """
    Carica le regole di classificazione dal file JSON.
    Compila le regex una volta sola e le mette in cache.
    """
    global _REGOLE_CACHE
    if _REGOLE_CACHE is not None:
        return _REGOLE_CACHE

    path = rules_path or _DEFAULT_RULES_PATH
    with open(path, "r", encoding="utf-8") as f:
        raw = json.load(f)

    # Compila le regex delle keywords
    for blk_name, regole_list in raw.items():
        for regola in regole_list:
            regola["_compiled"] = [
                re.compile(kw, re.I) for kw in regola.get("keywords_presenti", [])
            ]

    _REGOLE_CACHE = raw
    return _REGOLE_CACHE


def classifica_blocco(blk_name: str, testo: str) -> Tuple[str, str]:
    """
    Ritorna (convertibile, affidabilita') applicando le regole
    caricate dal file JSON esterno.

    Logiche supportate:
      - "ANY"         : matcha se ALMENO UNA keyword e' presente nel testo
      - "ALL"         : matcha se TUTTE le keywords sono presenti
      - "COUNT_GT_15" : matcha se la prima keyword ha piu' di 15 occorrenze
      - "DEFAULT"     : matcha sempre (fallback, deve essere l'ultima regola)
    """
    regole = carica_regole()
    regole_blk = regole.get(blk_name, [])

    for regola in regole_blk:
        logica = regola.get("logica", "DEFAULT")
        compiled: List[re.Pattern] = regola.get("_compiled", [])

        matched = False

        if logica == "DEFAULT":
            matched = True

        elif logica == "ANY":
            matched = any(r.search(testo) for r in compiled)

        elif logica == "ALL":
            matched = all(r.search(testo) for r in compiled)

        elif logica.startswith("COUNT_GT_"):
            # Es: "COUNT_GT_15" → soglia = 15
            soglia = int(logica.split("_")[-1])
            if compiled:
                count = len(compiled[0].findall(testo))
                matched = count > soglia

        if matched:
            return (regola["convertibile"], regola["affidabilita"])

    # Fallback se nessuna regola definita per questo blk_name
    return ("NO", "BASSA")


def build_regole_dataframe(rules_path: Path | None = None) -> pd.DataFrame:
    """
    Costruisce un DataFrame leggibile con tutte le regole di classificazione
    per il sheet Excel 'Regole Classificazione'.
    """
    regole = carica_regole(rules_path)
    rows = []
    ordine = 0
    for blk_name, regole_list in regole.items():
        macro_cat = MACRO_CATEGORIE.get(blk_name, blk_name.upper())
        for regola in regole_list:
            ordine += 1
            rows.append({
                "ordine": ordine,
                "blk_name": blk_name,
                "macro_categoria": macro_cat,
                "nome_regola": regola["nome_regola"],
                "logica": regola["logica"],
                "keywords": ", ".join(regola.get("keywords_presenti", [])) or "(nessuna)",
                "convertibile": regola["convertibile"],
                "affidabilita": regola["affidabilita"],
            })
    return pd.DataFrame(rows)


# ═══════════════════════════════════════════════════════════════════
# PARSER CON GERARCHIA (Passo 2)
# ═══════════════════════════════════════════════════════════════════
def parse_sas_blocks_tracked(
    path: str | Path,
    *,
    deep: bool = True,
    max_depth: int = 50,
    start_from: int = 1,
) -> Dict[str, Dict[str, object]]:
    """
    Parser iterativo (deep-first) con tracciatura gerarchica.

    Ogni blocco ha:
      - tipo, macro_categoria, linea_start, linea_stop, testo, chiuso
      - parent_uid (None per top-level)
      - depth (0 per top-level)
    """
    p = Path(path)
    file_id = hashlib.blake2b(str(p.resolve()).encode("utf-8"), digest_size=8).hexdigest()

    text = p.read_text(encoding="utf-8", errors="ignore")

    # Taglio opzionale mantenendo offset corretto
    if start_from and start_from > 1:
        all_lines = text.splitlines(keepends=True)
        text = "".join(all_lines[start_from - 1:])
        base_offset = start_from - 1
    else:
        base_offset = 0

    initial_raw: List[Tuple[str, bool]] = preprocess_lines_keep_comments(text)

    # ── work item = (raw_lines_segment, offset_line, depth, parent_uid) ──
    work = deque([(initial_raw, base_offset, 0, None)])

    blocks: Dict[str, Dict[str, object]] = {}
    counter = 1

    while work:
        raw_lines, offset, depth, parent_uid = work.popleft()
        i = 0
        n = len(raw_lines)

        while i < n:
            line, is_comment = raw_lines[i]
            if is_comment:
                i += 1
                continue

            matched = False

            for blk_name, spec in BLOCKS.items():
                start_re: re.Pattern = spec["start"]
                if start_re.match(line):
                    matched = True
                    tipo_label = spec["key_fn"](line.strip(), counter)
                    end_re: re.Pattern = spec["end"]
                    macro_cat = MACRO_CATEGORIE.get(blk_name, "SCONOSCIUTO")

                    start_idx = i
                    start_line_no = offset + start_idx + 1  # 1-based

                    buf: List[str] = [line]
                    i += 1

                    # ── %include chiuso su riga singola con ';' ──
                    if blk_name == "include" and line.rstrip().endswith(";"):
                        stop_line_no = start_line_no
                        block_text = "".join(buf)
                        conv, aff = classifica_blocco(blk_name, block_text)
                        uid = _stable_uid(file_id, blk_name, start_line_no, block_text)
                        blocks[uid] = {
                            "tipo": tipo_label,
                            "macro_categoria": macro_cat,
                            "linea_start": start_line_no,
                            "linea_stop": stop_line_no,
                            "testo": block_text,
                            "chiuso": True,
                            "parent_uid": parent_uid,
                            "depth": depth,
                            "convertibile_auto": conv,
                            "affidabilita_auto": aff,
                        }
                        counter += 1
                        break

                    # ── %run_macro chiuso su riga singola con ';' ──
                    if blk_name == "run_macro" and line.rstrip().endswith(";"):
                        stop_line_no = start_line_no
                        block_text = "".join(buf)
                        conv, aff = classifica_blocco(blk_name, block_text)
                        uid = _stable_uid(file_id, blk_name, start_line_no, block_text)
                        blocks[uid] = {
                            "tipo": tipo_label,
                            "macro_categoria": macro_cat,
                            "linea_start": start_line_no,
                            "linea_stop": stop_line_no,
                            "testo": block_text,
                            "chiuso": True,
                            "parent_uid": parent_uid,
                            "depth": depth,
                            "convertibile_auto": conv,
                            "affidabilita_auto": aff,
                        }
                        counter += 1
                        break

                    # ── ricerca end (bilanciata opzionale) ──
                    end_found = False
                    stop_line_no = None

                    balanced = bool(spec.get("balanced"))
                    nest_starts: List[re.Pattern] = spec.get("nest_starts", []) if balanced else []
                    end_token: re.Pattern = spec.get("end_token", end_re) if balanced else end_re

                    if balanced:
                        depth_if = 1
                        while i < n:
                            l, c = raw_lines[i]
                            buf.append(l)
                            if not c:
                                if any(r.match(l) for r in nest_starts):
                                    depth_if += 1
                                elif end_token.match(l):
                                    depth_if -= 1
                                    if depth_if == 0:
                                        end_found = True
                                        stop_line_no = offset + i + 1
                                        break
                            i += 1
                    else:
                        while i < n:
                            l, c = raw_lines[i]
                            buf.append(l)
                            if not c and end_re.match(l):
                                end_found = True
                                stop_line_no = offset + i + 1
                                break
                            i += 1

                    # ── salva blocco ──
                    block_text = "".join(buf)
                    conv, aff = classifica_blocco(blk_name, block_text)
                    uid = _stable_uid(file_id, blk_name, start_line_no, block_text)
                    blocks[uid] = {
                        "tipo": tipo_label,
                        "macro_categoria": macro_cat,
                        "linea_start": start_line_no,
                        "linea_stop": stop_line_no,
                        "testo": block_text,
                        "chiuso": bool(end_found),
                        "parent_uid": parent_uid,
                        "depth": depth,
                        "convertibile_auto": conv,
                        "affidabilita_auto": aff,
                    }
                    counter += 1

                    # ── deep push con parent_uid ──
                    if deep and depth < max_depth:
                        if end_found and len(buf) >= 2:
                            inner_text = "".join(buf[1:-1])
                            inner_offset = start_line_no
                        else:
                            inner_text = "".join(buf[1:])
                            inner_offset = start_line_no

                        if inner_text.strip():
                            inner_raw = preprocess_lines_keep_comments(inner_text)
                            work.appendleft((inner_raw, inner_offset, depth + 1, uid))

                    break  # blocco trovato

            if not matched:
                i += 1

    return blocks


# ═══════════════════════════════════════════════════════════════════
# COSTRUZIONE DATAFRAME (Passo 5 + 6)
# ═══════════════════════════════════════════════════════════════════
def build_dataframe(
    blocks: Dict[str, Dict[str, object]],
    pyspark_map: Dict[str, str] | None = None,
) -> pd.DataFrame:
    """
    Costruisce il DataFrame con tutte le colonne richieste.

    pyspark_map : dizionario opzionale {uid -> codice_pyspark}.
                  Se fornito, aggiunge la colonna 'codice_pyspark' all'Excel
                  per visualizzare la conversione blocco per blocco.
    """
    if not blocks:
        return pd.DataFrame()

    rows = []
    for uid, blk in blocks.items():
        rows.append({
            "uid": uid,
            "macro_categoria": blk["macro_categoria"],
            "tipo": blk["tipo"],
            "linea_start": blk["linea_start"],
            "linea_stop": blk["linea_stop"],
            "chiuso": blk["chiuso"],
            "parent_uid": blk.get("parent_uid"),
            "depth": blk.get("depth", 0),
            "children_uids": "",
            "convertibile_auto": blk["convertibile_auto"],
            "affidabilita_auto": blk["affidabilita_auto"],
            "convertibile_manuale": "",
            "affidabilita_manuale": "",
            "note": "",
            "testo": blk["testo"],
            # Colonna conversione PySpark (popolata se pyspark_map fornito)
            "codice_pyspark": (pyspark_map or {}).get(uid, ""),
        })

    df = pd.DataFrame(rows)
    df = df.set_index("uid", drop=False)

    # ── Calcolo children_uids ──
    for uid in df.index:
        children = df[df["parent_uid"] == uid].index.tolist()
        df.at[uid, "children_uids"] = ",".join(children) if children else ""

    return df


# ═══════════════════════════════════════════════════════════════════
# STATISTICHE RIASSUNTIVE
# ═══════════════════════════════════════════════════════════════════
def build_stats(df: pd.DataFrame) -> pd.DataFrame:
    """
    Crea un DataFrame di statistiche riassuntive per la sheet Excel / UI web.

    Garantisce che TUTTI i blocchi compaiano nelle righe di categoria:
    - usa value_counts(dropna=False) per includere categorie null/vuote
    - sostituisce i valori null/vuoti con "(non classificato)"
    - il TOTALE corrisponde sempre alla somma dei singoli tipi
    """
    if df.empty:
        return pd.DataFrame()

    # Normalizza le categorie mancanti
    df = df.copy()
    df["macro_categoria"] = (
        df["macro_categoria"].fillna("(non classificato)").replace("", "(non classificato)")
    )

    stats_rows = []

    # Conteggio per macro_categoria — dropna=False include anche categorie vuote
    cat_counts = df["macro_categoria"].value_counts(dropna=False)
    for cat, count in cat_counts.items():
        sub = df[df["macro_categoria"] == cat]
        n_conv_si = (sub["convertibile_auto"] == "SI").sum()
        n_conv_no = (sub["convertibile_auto"] == "NO").sum()
        n_alta  = (sub["affidabilita_auto"] == "ALTA").sum()
        n_media = (sub["affidabilita_auto"] == "MEDIA").sum()
        n_bassa = (sub["affidabilita_auto"] == "BASSA").sum()
        stats_rows.append({
            "macro_categoria":    str(cat) if cat is not None else "(non classificato)",
            "totale_blocchi":     int(count),
            "convertibile_SI":   int(n_conv_si),
            "convertibile_NO":   int(n_conv_no),
            "affidabilita_ALTA":  int(n_alta),
            "affidabilita_MEDIA": int(n_media),
            "affidabilita_BASSA": int(n_bassa),
        })

    # Riga TOTALE — calcolata sulla somma delle righe precedenti per coerenza
    total_blocchi = sum(r["totale_blocchi"]   for r in stats_rows)
    total_si      = sum(r["convertibile_SI"]  for r in stats_rows)
    total_no      = sum(r["convertibile_NO"]  for r in stats_rows)
    total_alta    = sum(r["affidabilita_ALTA"]  for r in stats_rows)
    total_media   = sum(r["affidabilita_MEDIA"] for r in stats_rows)
    total_bassa   = sum(r["affidabilita_BASSA"] for r in stats_rows)

    stats_rows.append({
        "macro_categoria":    "TOTALE",
        "totale_blocchi":     total_blocchi,
        "convertibile_SI":    total_si,
        "convertibile_NO":    total_no,
        "affidabilita_ALTA":  total_alta,
        "affidabilita_MEDIA": total_media,
        "affidabilita_BASSA": total_bassa,
    })

    return pd.DataFrame(stats_rows)


# ═══════════════════════════════════════════════════════════════════
# EXPORT EXCEL (Passo 7)
# ═══════════════════════════════════════════════════════════════════
def export_excel(df: pd.DataFrame, stats_df: pd.DataFrame, output_path: Path) -> None:
    """Salva il DataFrame in Excel con formattazione condizionale."""
    regole_df = build_regole_dataframe()

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        # Sheet principale
        df.to_excel(writer, sheet_name="Blocchi", index=False)
        # Sheet statistiche
        stats_df.to_excel(writer, sheet_name="Statistiche", index=False)
        # Sheet regole di classificazione
        regole_df.to_excel(writer, sheet_name="Regole Classificazione", index=False)

    # ── Formattazione con openpyxl ──
    from openpyxl import load_workbook
    from openpyxl.styles import PatternFill, Alignment, Font
    from openpyxl.utils import get_column_letter

    wb = load_workbook(output_path)

    # ── Sheet "Blocchi" ──
    ws = wb["Blocchi"]

    # Colori per affidabilita
    fill_alta  = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")  # verde
    fill_media = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")  # giallo
    fill_bassa = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")  # rosso

    # Trova indice colonne
    headers = [cell.value for cell in ws[1]]
    col_aff     = headers.index("affidabilita_auto") + 1 if "affidabilita_auto"  in headers else None
    col_conv    = headers.index("convertibile_auto") + 1 if "convertibile_auto"  in headers else None
    col_testo   = headers.index("testo")             + 1 if "testo"              in headers else None
    col_pyspark = headers.index("codice_pyspark")    + 1 if "codice_pyspark"     in headers else None

    # Font monospace per le celle codice
    font_code  = Font(name="Courier New", size=9)
    fill_py_si = PatternFill(start_color="DDEEFF", end_color="DDEEFF", fill_type="solid")  # azzurro chiaro
    fill_py_no = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")  # grigio chiaro

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        # Colore affidabilita_auto
        if col_aff:
            cell = row[col_aff - 1]
            if cell.value == "ALTA":
                cell.fill = fill_alta
            elif cell.value == "MEDIA":
                cell.fill = fill_media
            elif cell.value == "BASSA":
                cell.fill = fill_bassa

        # Colore convertibile_auto
        if col_conv:
            cell = row[col_conv - 1]
            if cell.value == "SI":
                cell.fill = fill_alta
            elif cell.value == "NO":
                cell.fill = fill_bassa

        # Wrap text colonna testo SAS
        if col_testo:
            row[col_testo - 1].alignment = Alignment(wrap_text=True, vertical="top")

        # Colonna codice_pyspark: font monospace + wrap + colore
        if col_pyspark:
            cell_py = row[col_pyspark - 1]
            cell_py.font      = font_code
            cell_py.alignment = Alignment(wrap_text=True, vertical="top")
            # Colore di sfondo in base al contenuto (TODO = grigio, codice = azzurro)
            val = str(cell_py.value or "")
            if val.strip().startswith("# TODO") or "revisione manuale" in val:
                cell_py.fill = fill_py_no
            elif val.strip():
                cell_py.fill = fill_py_si

    # Header bold
    header_font = Font(bold=True)
    for cell in ws[1]:
        cell.font = header_font

    # Larghezze colonne ragionevoli
    for col_idx, header in enumerate(headers, 1):
        col_letter = get_column_letter(col_idx)
        if header == "testo":
            ws.column_dimensions[col_letter].width = 70
        elif header == "codice_pyspark":
            ws.column_dimensions[col_letter].width = 80
        elif header in ("uid", "parent_uid", "children_uids"):
            ws.column_dimensions[col_letter].width = 40
        elif header == "tipo":
            ws.column_dimensions[col_letter].width = 25
        else:
            ws.column_dimensions[col_letter].width = 18

    # Filtri automatici
    ws.auto_filter.ref = ws.dimensions

    # ── Sheet "Statistiche" ──
    ws_stats = wb["Statistiche"]
    for cell in ws_stats[1]:
        cell.font = header_font
    for col_idx in range(1, ws_stats.max_column + 1):
        ws_stats.column_dimensions[get_column_letter(col_idx)].width = 20

    # ── Sheet "Regole Classificazione" ──
    ws_regole = wb["Regole Classificazione"]
    for cell in ws_regole[1]:
        cell.font = header_font

    # Colori per convertibile/affidabilita nelle regole
    regole_headers = [cell.value for cell in ws_regole[1]]
    col_reg_conv = regole_headers.index("convertibile") + 1 if "convertibile" in regole_headers else None
    col_reg_aff = regole_headers.index("affidabilita") + 1 if "affidabilita" in regole_headers else None

    for row in ws_regole.iter_rows(min_row=2, max_row=ws_regole.max_row):
        if col_reg_conv:
            cell = row[col_reg_conv - 1]
            if cell.value == "SI":
                cell.fill = fill_alta
            elif cell.value == "NO":
                cell.fill = fill_bassa
        if col_reg_aff:
            cell = row[col_reg_aff - 1]
            if cell.value == "ALTA":
                cell.fill = fill_alta
            elif cell.value == "MEDIA":
                cell.fill = fill_media
            elif cell.value == "BASSA":
                cell.fill = fill_bassa

    # Larghezze colonne regole
    regole_widths = {
        "ordine": 8, "blk_name": 15, "macro_categoria": 18,
        "nome_regola": 50, "logica": 15, "keywords": 55,
        "convertibile": 14, "affidabilita": 14,
    }
    for col_idx, header in enumerate(regole_headers, 1):
        ws_regole.column_dimensions[get_column_letter(col_idx)].width = regole_widths.get(header, 15)

    ws_regole.auto_filter.ref = ws_regole.dimensions

    # ── Tabella Priorita' di Conversione (sotto le regole) ──
    # 2 righe vuote di separazione
    start_row = ws_regole.max_row + 3

    # Titolo
    title_cell = ws_regole.cell(row=start_row, column=1,
                                value="PRIORITA' DI CONVERSIONE")
    title_cell.font = Font(bold=True, size=13)

    # Header tabella priorita'
    priorita_headers = ["Priorita'", "Convertibile", "Affidabilita'",
                        "Azione", "Descrizione"]
    for col_idx, h in enumerate(priorita_headers, 1):
        cell = ws_regole.cell(row=start_row + 1, column=col_idx, value=h)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2",
                                fill_type="solid")

    # Righe della tabella
    priorita_data = [
        (1, "SI", "ALTA",
         "Conversione AUTOMATICA",
         "Il mapping sintattico SAS->PySpark e' sufficiente. "
         "Nessun intervento manuale richiesto."),
        (2, "SI", "MEDIA",
         "Conversione SEMI-AUTOMATICA",
         "Il motore propone una conversione, ma serve una revisione "
         "umana per validare la logica (es. MERGE, IF/ELSE, PROC DATASETS)."),
        (3, "SI", "BASSA",
         "Conversione MANUALE ASSISTITA",
         "Il motore propone un'abbozzo, ma il contenuto e' sconosciuto "
         "o complesso (es. macro call, macro con molte variabili). "
         "Lo sviluppatore deve completare la conversione."),
        (4, "NO", "BASSA",
         "RISCRITTURA COMPLETA",
         "Blocco non convertibile automaticamente (es. HASH, ARRAY, RETAIN, "
         "file esterni %INCLUDE). Lo sviluppatore deve riscrivere "
         "interamente la logica in PySpark."),
    ]

    for row_offset, (pri, conv, aff, azione, desc) in enumerate(priorita_data):
        r = start_row + 2 + row_offset
        ws_regole.cell(row=r, column=1, value=pri)
        c_conv = ws_regole.cell(row=r, column=2, value=conv)
        c_aff  = ws_regole.cell(row=r, column=3, value=aff)
        ws_regole.cell(row=r, column=4, value=azione)
        c_desc = ws_regole.cell(row=r, column=5, value=desc)
        c_desc.alignment = Alignment(wrap_text=True, vertical="top")

        # Colori coerenti
        if conv == "SI":
            c_conv.fill = fill_alta
        else:
            c_conv.fill = fill_bassa
        if aff == "ALTA":
            c_aff.fill = fill_alta
        elif aff == "MEDIA":
            c_aff.fill = fill_media
        else:
            c_aff.fill = fill_bassa

    # Larghezze colonne per la tabella priorita' (riusa le stesse colonne)
    ws_regole.column_dimensions[get_column_letter(4)].width = 35
    ws_regole.column_dimensions[get_column_letter(5)].width = 70

    wb.save(output_path)
    print(f"Excel salvato: {output_path}")


# ═══════════════════════════════════════════════════════════════════
# EXPORT JSON AD ALBERO (Passo 7)
# ═══════════════════════════════════════════════════════════════════
def _build_tree(blocks: Dict[str, Dict[str, object]]) -> list:
    """Costruisce struttura ad albero annidato per il JSON."""
    # Indice: uid -> lista figli
    children_map: Dict[Optional[str], list] = {}
    for uid, blk in blocks.items():
        p = blk.get("parent_uid")
        children_map.setdefault(p, []).append(uid)

    def _node(uid: str) -> dict:
        blk = blocks[uid]
        node = {
            "uid": uid,
            "tipo": blk["tipo"],
            "macro_categoria": blk["macro_categoria"],
            "linea_start": blk["linea_start"],
            "linea_stop": blk["linea_stop"],
            "chiuso": blk["chiuso"],
            "depth": blk.get("depth", 0),
            "convertibile_auto": blk["convertibile_auto"],
            "affidabilita_auto": blk["affidabilita_auto"],
            "testo": blk["testo"],
        }
        kids = children_map.get(uid, [])
        if kids:
            node["children"] = [_node(c) for c in kids]
        return node

    # Nodi root = quelli con parent_uid == None
    roots = children_map.get(None, [])
    return [_node(r) for r in roots]


def export_json(blocks: Dict[str, Dict[str, object]], output_path: Path) -> None:
    """Salva la struttura ad albero in JSON."""
    tree = _build_tree(blocks)
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(tree, f, ensure_ascii=False, indent=2, default=str)
    print(f"JSON salvato: {output_path}")


# ═══════════════════════════════════════════════════════════════════
# MAIN (Passo 8)
# ═══════════════════════════════════════════════════════════════════
def main():
    parser = argparse.ArgumentParser(
        description="Parser SAS con tracciatura gerarchica dei blocchi"
    )
    parser.add_argument("sas_file", help="Percorso file SAS da analizzare")
    parser.add_argument("--output-dir", default=".", help="Directory output (default: corrente)")
    parser.add_argument("--rules", default=None, help="Percorso file JSON regole classificazione (default: regole_classificazione.json)")
    parser.add_argument("--no-excel", action="store_true", help="Non generare file Excel")
    parser.add_argument("--no-json", action="store_true", help="Non generare file JSON")
    args = parser.parse_args()

    sas_path = Path(args.sas_file)
    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    if not sas_path.exists():
        print(f"ERRORE: File non trovato: {sas_path}")
        return

    # ── Carica regole ──
    rules_path = Path(args.rules) if args.rules else None
    carica_regole(rules_path)

    # ── Parse ──
    print(f"Analisi di: {sas_path}")
    blocks = parse_sas_blocks_tracked(sas_path)
    print(f"Blocchi trovati: {len(blocks)}")


    for k, v in blocks.items():
        print(f"\n=== {k} ===\n{v['testo']}")
      
    # ── DataFrame ──
    df = build_dataframe(blocks)

    # ── Stampa a console ──
    pd.set_option("display.max_columns", None)
    pd.set_option("display.width", 200)
    pd.set_option("display.max_colwidth", 60)

    cols_display = [c for c in df.columns if c != "testo"]
    print("\n" + df[cols_display].to_string())

    # ── Statistiche ──
    stats_df = build_stats(df)
    print("\n=== STATISTICHE ===")
    print(stats_df.to_string(index=False))

    # ── Timestamp per nomi file: ANNO-MESE-GIORNO-ORA-MINUTO-SECONDO-MILLISECONDO ──
    now = datetime.now()
    ts = now.strftime("%Y-%m-%d-%H-%M-%S") + f"-{now.microsecond // 1000:03d}"

    # ── Export Excel ──
    if not args.no_excel:
        excel_path = output_dir / f"{ts}.xlsx"
        export_excel(df, stats_df, excel_path)

    # ── Export JSON ──
    if not args.no_json:
        json_path = output_dir / f"{ts}.json"
        export_json(blocks, json_path)

    print("\nDone.")


if __name__ == "__main__":
    main()
